"""Build the running Basin Bench deck. Re-run after every milestone; slides/basin_bench.pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
TEAL = RGBColor(0x1C, 0x72, 0x93)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MIST = RGBColor(0xE6, 0xEF, 0xF5)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb


def number(slide, n, color):
    text(slide, 12.3, 6.9, 0.8, 0.4, str(n), 18, color=color, align=PP_ALIGN.RIGHT)


def fill(slide, rgb):
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb


# Slide 1: the question
s = prs.slides.add_slide(blank); fill(s, NAVY)
text(s, 0.8, 0.9, 11.7, 1.2, "BASIN BENCH", 40, bold=True, color=RGBColor(0x9F, 0xD3, 0xE8))
text(s, 0.8, 2.2, 11.7, 3.0, "Does physics make a machine learning model better at predicting rivers?",
     60, bold=True, color=WHITE)
text(s, 0.8, 5.6, 11.7, 0.8, "671 US rivers.  Three models.  One fair test.", 32, color=RGBColor(0x9F, 0xD3, 0xE8))
number(s, 1, RGBColor(0x9F, 0xD3, 0xE8))
s.notes_slide.notes_text_frame.text = (
    "Read the question slowly. Pause. Rivers are hard to predict. We ask one thing: "
    "when a learning model is handed what a physics model knows, does it predict better.")

# Slide 2: three models, same rivers, one score
s = prs.slides.add_slide(blank); fill(s, WHITE)
text(s, 0.8, 0.6, 11.7, 1.4, "Three models. The same rivers. One score each.", 48, bold=True, color=NAVY)
cards = [("PHYSICS", "GR4J", "Keeps account of water in the soil"),
         ("LEARNER", "XGBoost", "Finds patterns in past weather"),
         ("BOTH", "Hybrid", "The learner, handed the physics account")]
x0, cw, gap = 0.8, 3.7, 0.3
for i, (kind, name, line) in enumerate(cards):
    x = x0 + i * (cw + gap)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.3), Inches(cw), Inches(3.2))
    box.fill.solid(); box.fill.fore_color.rgb = MIST if i < 2 else TEAL
    box.line.fill.background(); box.adjustments[0] = 0.08
    c = NAVY if i < 2 else WHITE
    text(s, x + 0.3, 2.55, cw - 0.6, 0.5, kind, 20, bold=True, color=c)
    text(s, x + 0.3, 3.05, cw - 0.6, 1.0, name, 44, bold=True, color=c)
    text(s, x + 0.3, 4.2, cw - 0.6, 1.2, line, 22, color=c)
text(s, 0.8, 5.9, 11.7, 0.9, "Every river gives every model one score. We compare the scores river by river.",
     28, bold=True, color=TEAL)
number(s, 2, NAVY)
s.notes_slide.notes_text_frame.text = (
    "Three contestants. The physics model keeps a running account of soil water. The learner only sees "
    "patterns. The hybrid is the learner handed the physics account. Same rivers, same years, one score each. "
    "Pause. The comparison is river by river, which is what makes it fair.")

prs.save("slides/basin_bench.pptx")
print("saved slides/basin_bench.pptx")
